"""Endpoint wiring for compound pre-flight reconciliation in
UploadExperimentValidateView. The classification logic itself is covered by
test_compound_reconciliation; here we only prove the view accepts the compound
list and surfaces conflicts in the response."""

# pylint: disable=redefined-outer-name,unused-argument

import base64
import io

import pytest
from openpyxl import load_workbook

from viewer.compound_reconciliation import inchi_key_for_smiles
from viewer.models import Compound, Project
from viewer.tests.curation_sheet_helpers import SHEET, edit_group_row, existing_row_ids

VALIDATE_URL = "/api/validate_target_experiments/"
TAS = "lb00001-1"
ETHANOL = "CCO"


@pytest.fixture
def bypass_auth_and_versions(settings, monkeypatch):
    """Skip upload auth and the (unrelated) version checks so the test isolates
    the compound-reconciliation branch of the view."""
    settings.AUTHENTICATE_UPLOAD = False
    monkeypatch.setattr("viewer.views.split_version", lambda _v: (1, 0))
    monkeypatch.setattr(
        "viewer.views.validate_data_version", lambda *a, **k: (True, "data ok")
    )
    monkeypatch.setattr(
        "viewer.views.validate_upload_version", lambda *a, **k: (True, "upload ok")
    )


def _payload(compounds):
    return {
        "target_access_string": TAS,
        "target_name": "Mpro",
        "data_version": "1.0",
        "upload_version": "1",
        "compounds": compounds,
    }


@pytest.mark.django_db
def test_validate_reports_compound_conflict(api_client, bypass_auth_and_versions):
    project = Project.objects.create(title=TAS)
    Compound.objects.create(
        project=project,
        smiles=ETHANOL,
        inchi="InChI=CCO",
        inchi_key=inchi_key_for_smiles(ETHANOL),
        compound_code="OLD",
    )
    response = api_client.post(
        VALIDATE_URL,
        _payload([{"smiles": ETHANOL, "compound_code": "NEW"}]),
        format="json",
    )
    assert response.status_code == 200
    assert response.data["success"] is False
    conflicts = response.data["compound_conflicts"]
    assert len(conflicts) == 1
    assert conflicts[0]["status"] == "conflict"
    assert conflicts[0]["conflicts"]["compound_code"] == {
        "existing": "OLD",
        "incoming": "NEW",
    }


@pytest.mark.django_db
def test_validate_clean_when_no_conflict(api_client, bypass_auth_and_versions):
    project = Project.objects.create(title=TAS)
    Compound.objects.create(
        project=project,
        smiles=ETHANOL,
        inchi="InChI=CCO",
        inchi_key=inchi_key_for_smiles(ETHANOL),
        compound_code="SAME",
    )
    response = api_client.post(
        VALIDATE_URL,
        _payload([{"smiles": ETHANOL, "compound_code": "SAME"}]),
        format="json",
    )
    assert response.status_code == 200
    assert response.data["success"] is True
    assert "compound_conflicts" not in response.data


@pytest.mark.django_db
def test_validate_conflict_returns_curation_spreadsheet(
    api_client, bypass_auth_and_versions
):
    project = Project.objects.create(title=TAS)
    Compound.objects.create(
        project=project,
        smiles=ETHANOL,
        inchi="InChI=CCO",
        inchi_key=inchi_key_for_smiles(ETHANOL),
        compound_code="OLD",
    )
    response = api_client.post(
        VALIDATE_URL,
        _payload([{"smiles": ETHANOL, "compound_code": "NEW"}]),
        format="json",
    )
    assert response.data["success"] is False
    assert response.data["curation_filename"].endswith(".xlsx")
    # the returned base64 decodes to a real workbook carrying our conflict sheet
    data = base64.b64decode(response.data["curation_file"])
    wb = load_workbook(io.BytesIO(data))
    assert SHEET in wb.sheetnames


@pytest.mark.django_db
def test_validate_with_completed_curation_resolves(
    api_client, bypass_auth_and_versions
):
    project = Project.objects.create(title=TAS)
    Compound.objects.create(
        project=project,
        smiles=ETHANOL,
        inchi="InChI=CCO",
        inchi_key=inchi_key_for_smiles(ETHANOL),
        compound_code="OLD",
    )
    compounds = [{"smiles": ETHANOL, "compound_code": "NEW"}]

    # first pass: get the generated spreadsheet
    first = api_client.post(VALIDATE_URL, _payload(compounds), format="json")
    assert first.data["success"] is False
    data = base64.b64decode(first.data["curation_file"])

    # user resolves the single conflict by keeping the existing compound
    existing_id = existing_row_ids(data)[0]
    data = edit_group_row(data, existing_id, action="KEEP")

    payload = _payload(compounds)
    payload["curation_file"] = base64.b64encode(data).decode("ascii")
    second = api_client.post(VALIDATE_URL, payload, format="json")
    assert second.data["success"] is True
    assert "compound_conflicts" not in second.data


@pytest.mark.django_db
def test_validate_without_compounds_is_unaffected(api_client, bypass_auth_and_versions):
    """Legacy callers that don't send compounds still validate normally."""
    payload = _payload([])
    payload.pop("compounds")
    response = api_client.post(VALIDATE_URL, payload, format="json")
    assert response.status_code == 200
    assert response.data["success"] is True
    assert "compound_conflicts" not in response.data
