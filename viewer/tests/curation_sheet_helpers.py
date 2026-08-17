"""Helpers for editing a generated curation workbook the way a user would.

The "Incoming compounds" sheet is a stack of per-compound groups, each preceded
by its own header row. These helpers locate cells by column *name* rather than
number, so they keep working as CONTENT_FIELDS (and therefore the sheet's width)
changes.
"""

import io

from openpyxl import load_workbook

SHEET = "Incoming compounds"


def edit_group_row(data: bytes, marker, **cell_values) -> bytes:
    """Set named columns on every row whose id column equals ``marker``.

    ``marker`` is "Incoming", "Merge", or an existing compound's database id.
    Returns the re-saved workbook bytes.
    """
    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET]
    cols: dict = {}
    for row in ws.iter_rows():
        first = row[0].value
        if first == "id":
            cols = {c.value: c.column for c in row if c.value}
            continue
        if str(first).strip().lower() == str(marker).strip().lower():
            for name, value in cell_values.items():
                ws.cell(row[0].row, cols[name], value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def group_headers(data: bytes):
    """Column names of the first group's header row."""
    ws = load_workbook(io.BytesIO(data))[SHEET]
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "id":
            return [v for v in row if v]
    return []


def existing_row_ids(data: bytes):
    """Database ids appearing in the id column of existing-duplicate rows."""
    ws = load_workbook(io.BytesIO(data))[SHEET]
    ids = []
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] in (None, "", "id"):
            continue
        if str(row[0]).strip().lower() in ("incoming", "merge"):
            continue
        ids.append(row[0])
    return ids
