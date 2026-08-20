"""Standalone (no-DB) tests for the hierarchical compound-group curation sheet.

These cover the build -> edit -> parse round trip of the "Incoming compounds"
sheet, and in particular that its column layout is derived from CONTENT_FIELDS
on both sides rather than hardcoded. They need only openpyxl.
"""

import io

import pytest
from openpyxl import load_workbook

from viewer import compound_curation as cc
from viewer.compound_curation import (
    MERGE_ROW_ACTION,
    ROW_MARKER_INCOMING,
    CurationFormatError,
    build_curation_xlsx,
    curation_headers,
    parse_curation_xlsx,
    resolve_curation,
)
from viewer.tests.curation_sheet_helpers import INCOMING_ROW, MERGE_ROW, SHEET
from viewer.tests.curation_sheet_helpers import edit_group_row as _edit
from viewer.tests.curation_sheet_helpers import existing_row_ids
from viewer.tests.curation_sheet_helpers import group_headers as _headers_in

INCOMING = {
    "smiles": "CCO",
    "compound_code": "NEW-1",
    "ligand_name": "LIG-NEW",
    "modeled_smiles_canon": "CCO",
}
EXISTING = {
    "id": 42,
    "inchi_key": "AAA",
    "smiles": "CCO",
    "compound_code": "OLD-1",
    "ligand_name": "LIG-OLD",
    "modeled_smiles_canon": "CCO",
}
KEY = ("AAA", "CCO", "NEW-1")


def _conflict_payload():
    return [
        {
            "status": "conflict",
            "inchi_key": "AAA",
            "incoming": dict(INCOMING),
            "conflicts": {
                "compound_code": {"existing": "OLD-1", "incoming": "NEW-1"},
                "ligand_name": {"existing": "LIG-OLD", "incoming": "LIG-NEW"},
            },
            "existing": [dict(EXISTING)],
        }
    ]


def test_headers_follow_content_fields():
    headers = curation_headers()
    assert headers[0] == "id"
    assert headers[-1] == "action"
    # every content field appears exactly once, in a stable order
    for f in cc.CONTENT_FIELDS:
        assert headers.count(f) == 1
    assert _headers_in(build_curation_xlsx(_conflict_payload())) == headers


def test_untouched_sheet_round_trips_to_defaults():
    data = build_curation_xlsx(_conflict_payload())
    parsed = parse_curation_xlsx(data)

    # incoming row defaults to CREATE, existing duplicate to KEEP
    assert parsed.field_actions[KEY] == "CREATE"
    assert parsed.field_actions[(*KEY, 42)] == "KEEP"
    # merge row is blank until the user picks something
    assert ("MERGE", KEY) not in parsed.field_actions


def test_existing_row_is_keyed_by_incoming_identity_not_its_own():
    """The existing row carries OLD-1 in compound_code; the conflict is exactly
    that it differs from the incoming NEW-1. The decision must still be found
    under the incoming identity, which is what resolve_curation looks up."""
    data = build_curation_xlsx(_conflict_payload())
    parsed = parse_curation_xlsx(data)

    assert (*KEY, 42) in parsed.field_actions
    assert ("AAA", "CCO", "OLD-1", 42) not in parsed.field_actions


def test_delete_decision_reaches_the_resolver():
    payload = _conflict_payload()
    data = _edit(build_curation_xlsx(payload), 42, action="DELETE")
    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert plan.ok
    assert [(a.op, a.existing_id) for a in plan.actions] == [("create", None)]


def test_defaults_keep_existing_and_create_incoming():
    """KEEP on the existing row plus CREATE on the incoming row means exactly
    that: the existing compound is untouched and the incoming one is added."""
    payload = _conflict_payload()
    plan = resolve_curation(payload, parse_curation_xlsx(build_curation_xlsx(payload)))

    assert plan.ok
    assert [(a.op, a.existing_id) for a in plan.actions] == [("create", None)]
    assert plan.actions[0].superseded_ids == []


def test_incoming_folds_into_the_single_kept_row():
    """Taking CREATE off the incoming row is how the user says "this is the same
    compound as that one" -- then the kept existing row is reused."""
    payload = _conflict_payload()
    data = _edit(build_curation_xlsx(payload), INCOMING_ROW, action="KEEP")
    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert plan.ok
    assert [(a.op, a.existing_id) for a in plan.actions] == [("reuse", 42)]


def test_merge_row_is_read_back():
    payload = _conflict_payload()
    data = _edit(
        build_curation_xlsx(payload),
        MERGE_ROW,
        action="MERGE",
        compound_code="OLD-1",
        ligand_name="LIG-NEW",
    )
    parsed = parse_curation_xlsx(data)

    merged = parsed.field_actions[("MERGE", KEY)]
    # compound_code is a conflict field like any other and must survive the trip
    assert merged["compound_code"] == "OLD-1"
    assert merged["ligand_name"] == "LIG-NEW"


def test_merge_decision_reaches_the_resolver():
    payload = _conflict_payload()
    data = _edit(
        build_curation_xlsx(payload),
        MERGE_ROW,
        action="MERGE",
        compound_code="OLD-1",
        ligand_name="LIG-NEW",
    )
    data = _edit(data, 42, action="MERGE")
    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert plan.ok
    (action,) = plan.actions
    assert action.op == "update"
    assert action.existing_id == 42
    assert action.incoming["compound_code"] == "OLD-1"
    assert action.incoming["ligand_name"] == "LIG-NEW"


def test_blank_merge_cell_fails_rather_than_guessing():
    """A merge-row cell the sheet asked the user to decide must not fall back to
    the incoming value -- that would silently pick a winner on their behalf."""
    payload = _conflict_payload()
    # user chose MERGE but never picked a compound_code from the dropdown
    data = _edit(
        build_curation_xlsx(payload), MERGE_ROW, action="MERGE", ligand_name="LIG-OLD"
    )
    data = _edit(data, 42, action="MERGE")
    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert not plan.ok
    assert not plan.actions
    assert "compound_code" in plan.unresolved[0].reason


def test_bad_action_is_rejected():
    data = _edit(build_curation_xlsx(_conflict_payload()), 42, action="DESTROY")
    with pytest.raises(CurationFormatError):
        parse_curation_xlsx(data)


def test_layout_adapts_to_an_added_content_field(monkeypatch):
    """The whole point of the dynamic layout: adding a field to CONTENT_FIELDS
    widens the sheet and round-trips, with no column numbers to update."""
    extra = "assay_comment"
    content = cc.CONTENT_FIELDS + (extra,)
    monkeypatch.setattr(cc, "CONTENT_FIELDS", content)
    monkeypatch.setattr(
        cc,
        "CONFLICT_FIELDS",
        tuple(f for f in content if f not in cc.LOCKED_IDENTITY_FIELDS),
    )

    payload = _conflict_payload()
    payload[0]["incoming"][extra] = "clean"
    payload[0]["existing"][0][extra] = "precipitate"
    payload[0]["conflicts"][extra] = {
        "existing": "precipitate",
        "incoming": "clean",
    }

    data = build_curation_xlsx(payload)
    assert extra in _headers_in(data)

    data = _edit(data, MERGE_ROW, action="MERGE", **{extra: "precipitate"})
    parsed = parse_curation_xlsx(data)

    # decisions on the pre-existing columns still land correctly...
    assert parsed.field_actions[KEY] == "CREATE"
    assert parsed.field_actions[(*KEY, 42)] == "KEEP"
    # ...and the new column is read back rather than silently dropped
    assert parsed.field_actions[("MERGE", KEY)][extra] == "precipitate"


def test_group_without_incoming_row_is_rejected():
    data = build_curation_xlsx(_conflict_payload())
    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET]
    for row in ws.iter_rows():
        if str(row[0].value).strip().lower() == "incoming":
            ws.delete_rows(row[0].row)
            break
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(CurationFormatError):
        parse_curation_xlsx(buf.getvalue())


def _ambiguous_payload(ids=(39, 40)):
    """Two existing compounds, same structure, different compound_code -- the
    shape of the blocks in a real curation sheet."""
    return [
        {
            "status": "ambiguous",
            "inchi_key": "AAA",
            "incoming": dict(INCOMING),
            "conflicts": {},
            "existing": [
                {
                    "id": ex_id,
                    "inchi_key": "AAA",
                    "smiles": "CCO",
                    "compound_code": f"OLD-{ex_id}",
                    "ligand_name": "LIG-NEW",
                    "modeled_smiles_canon": "CCO",
                }
                for ex_id in ids
            ],
        }
    ]


def test_merge_with_all_actions_deleted_supersedes_the_old_rows():
    """Case 1: DELETE, DELETE, DELETE, MERGE with the merge row filled in."""
    payload = _ambiguous_payload()
    data = build_curation_xlsx(payload)
    for ex_id in (39, 40):
        data = _edit(data, ex_id, action="DELETE")
    data = _edit(data, "Incoming", action="DELETE")
    data = _edit(data, MERGE_ROW, action="MERGE", compound_code="CHOSEN")

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert plan.ok
    (action,) = plan.actions
    assert action.incoming["compound_code"] == "CHOSEN"
    # one compound survives, carrying the merge row's values; the other retires
    assert action.existing_id == 39
    assert action.superseded_ids == [40]


def test_merge_overrides_default_keep_and_create_actions():
    """Case 2: only MERGE selected, the other rows left at their defaults."""
    payload = _ambiguous_payload(ids=(46, 47))
    data = build_curation_xlsx(payload)
    data = _edit(data, MERGE_ROW, action="MERGE", compound_code="CHOSEN")

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert plan.ok
    # exactly one decision -- not a reuse from the defaulted KEEP *and* a create
    (action,) = plan.actions
    assert action.op == "update"
    assert action.incoming["compound_code"] == "CHOSEN"
    assert action.existing_id == 46
    assert action.superseded_ids == [47]


def test_merge_with_unfilled_conflict_fields_is_unresolved():
    """Case 3: MERGE selected but the conflicting cells left blank. This must
    fail the upload rather than quietly picking a winner."""
    payload = _ambiguous_payload(ids=(53, 54))
    # a second conflicting field, so both are left for the user to decide
    payload[0]["existing"][0]["modeled_smiles_soakdb"] = "CCO"
    payload[0]["existing"][1]["modeled_smiles_soakdb"] = "OCC"
    payload[0]["incoming"]["modeled_smiles_soakdb"] = "C-C-O"

    data = build_curation_xlsx(payload)
    data = _edit(data, MERGE_ROW, action="MERGE")  # no values chosen

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert not plan.ok
    assert not plan.actions
    reason = plan.unresolved[0].reason
    assert "compound_code" in reason and "modeled_smiles_soakdb" in reason


def test_merge_with_only_some_fields_filled_is_unresolved():
    payload = _ambiguous_payload(ids=(53, 54))
    payload[0]["existing"][0]["modeled_smiles_soakdb"] = "CCO"
    payload[0]["existing"][1]["modeled_smiles_soakdb"] = "OCC"
    payload[0]["incoming"]["modeled_smiles_soakdb"] = "C-C-O"

    data = build_curation_xlsx(payload)
    data = _edit(data, MERGE_ROW, action="MERGE", compound_code="CHOSEN")

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert not plan.ok
    assert "modeled_smiles_soakdb" in plan.unresolved[0].reason
    assert "compound_code" not in plan.unresolved[0].reason


RED = "FFFFC7CE"  # a decision we could not accept
YELLOW = "FFFFF3C4"  # editable, awaiting the user


def _cells_with_fill(ws, rgb):
    return [c for row in ws.iter_rows() for c in row if c.fill.start_color.rgb == rgb]


def _merge_row_cells(data):
    """(required cell, settled cell) from the merge row of a curation workbook."""
    ws = load_workbook(io.BytesIO(data))[SHEET]

    cols, merge_row = {}, None
    for row in ws.iter_rows():
        if row[0].value == "id":
            cols = {c.value: c.column for c in row if c.value}
        elif str(row[0].value).strip().lower() == MERGE_ROW.lower():
            merge_row = row[0].row

    return ws.cell(merge_row, cols["compound_code"]), ws.cell(
        merge_row, cols["ligand_name"]
    )


def test_required_merge_cells_are_marked_and_editable():
    """The cells the user must fill are unlocked and yellow, matching the action
    column; the rest are pre-filled and locked."""
    required, settled = _merge_row_cells(build_curation_xlsx(_ambiguous_payload()))

    assert required.protection.locked is False
    assert required.fill.start_color.rgb == YELLOW
    assert required.value in (None, "")
    assert settled.protection.locked is True
    assert settled.value == "LIG-NEW"


def test_reissued_workbook_flags_unresolved_decisions_in_red():
    """Red is reserved for a re-run: it means "we rejected this answer", so it
    must never appear on the workbook the user is first handed."""
    payload = _ambiguous_payload()

    first, _ = _merge_row_cells(build_curation_xlsx(payload))
    reissued, settled = _merge_row_cells(
        build_curation_xlsx(payload, error_cells={KEY: ["compound_code"]})
    )

    assert first.fill.start_color.rgb == YELLOW
    assert reissued.fill.start_color.rgb == RED
    # Only the outstanding decision changes colour; settled cells stay locked.
    assert reissued.protection.locked is False
    assert settled.protection.locked is True


def test_only_the_rejected_cell_turns_red():
    """The bug this guards: one bad answer must not redden every decision in the
    workbook. Two groups, one mistake -> exactly one red cell."""
    payload = _conflict_payload() + _ambiguous_payload()
    data = build_curation_xlsx(payload, error_cells={KEY: ["compound_code"]})
    ws = load_workbook(io.BytesIO(data))[SHEET]

    assert len(_cells_with_fill(ws, RED)) == 1
    assert _cells_with_fill(ws, YELLOW), "other groups must stay editable yellow"


def test_action_column_error_reddens_the_action_cells():
    """An empty field list means the fault is the action column itself."""
    payload = _ambiguous_payload()
    data = build_curation_xlsx(payload, error_cells={KEY: []})
    ws = load_workbook(io.BytesIO(data))[SHEET]

    action_col = curation_headers().index("action") + 1
    red = [c for c in _cells_with_fill(ws, RED) if c.column == action_col]

    # two existing rows + the incoming row; the merge row stays yellow because
    # choosing MERGE is one of the ways out.
    assert len(red) == 3
    assert not [c for c in _cells_with_fill(ws, RED) if c.column != action_col]


def test_untouched_sheet_keeps_existing_and_creates_incoming():
    """The sheet's defaults say KEEP/KEEP/CREATE, so that is what happens: both
    existing rows are left alone and the incoming compound is added as new. An
    untouched sheet must always be satisfiable, however many duplicates it holds.
    """
    payload = _ambiguous_payload()
    plan = resolve_curation(payload, parse_curation_xlsx(build_curation_xlsx(payload)))

    assert plan.ok, [u.reason for u in plan.unresolved]
    assert [a.op for a in plan.actions] == ["create"]
    # Nothing was retired - the existing rows are untouched, not superseded.
    assert plan.actions[0].superseded_ids == []


def test_folding_into_existing_needs_exactly_one_keep():
    """Only when the incoming compound is NOT being created does a second KEEP
    become a contradiction."""
    payload = _ambiguous_payload()
    data = _edit(build_curation_xlsx(payload), INCOMING_ROW, action="KEEP")

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert not plan.ok
    assert "exactly one KEEP" in plan.unresolved[0].reason
    # The action column is at fault, not any merge cell.
    assert plan.unresolved[0].fields == []
    assert plan.unresolved[0].key is not None


def test_deleting_every_row_creates_incoming_and_supersedes():
    payload = _ambiguous_payload()
    data = build_curation_xlsx(payload)
    for ex_id in (39, 40):
        data = _edit(data, ex_id, action="DELETE")

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert plan.ok
    (action,) = plan.actions
    assert action.op == "create"
    assert action.existing_id is None
    assert sorted(action.superseded_ids) == [39, 40]


def test_multiple_groups_are_kept_separate():
    payload = _conflict_payload()
    payload.append(
        {
            "status": "conflict",
            "inchi_key": "BBB",
            "incoming": {"smiles": "CCC", "compound_code": "NEW-2"},
            "conflicts": {"compound_code": {"existing": "OLD-2", "incoming": "NEW-2"}},
            "existing": [
                {
                    "id": 99,
                    "inchi_key": "BBB",
                    "smiles": "CCC",
                    "compound_code": "OLD-2",
                }
            ],
        }
    )
    data = _edit(build_curation_xlsx(payload), 99, action="DELETE")
    parsed = parse_curation_xlsx(data)

    assert parsed.field_actions[(*KEY, 42)] == "KEEP"
    assert parsed.field_actions[("BBB", "CCC", "NEW-2", 99)] == "DELETE"


def test_merge_row_is_labelled_merge_to():
    ws = load_workbook(io.BytesIO(build_curation_xlsx(_conflict_payload())))[SHEET]
    labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]

    assert MERGE_ROW == "Merge to"
    assert MERGE_ROW in labels


def test_an_unknown_row_marker_is_rejected_not_misread():
    """Strict vocabulary. An unrecognised marker used to fall through to the
    existing-duplicate branch, where the literal word became a compound id and
    its action was filed under it -- a decision no lookup could ever match.
    """
    data = build_curation_xlsx(_conflict_payload())
    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == MERGE_ROW:
            ws.cell(r, 1, "Merge")  # the label used by an older version
            ws.cell(r, curation_headers().index("action") + 1, MERGE_ROW_ACTION)
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(CurationFormatError, match="neither a compound id"):
        parse_curation_xlsx(buf.getvalue())


def _action_validations(data):
    """{first cell of the range: allowed-values formula} for the action column."""
    ws = load_workbook(io.BytesIO(data))[SHEET]
    action_col = curation_headers().index("action") + 1
    out = {}
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            cell = ws[rng.split(":")[0]]
            if cell.column == action_col:
                out[ws.cell(cell.row, 1).value] = dv.formula1
    return out


def test_each_row_kind_offers_only_its_own_actions():
    dvs = _action_validations(build_curation_xlsx(_ambiguous_payload()))

    assert dvs[39] == '"KEEP,DELETE"'
    assert dvs[40] == '"KEEP,DELETE"'
    assert dvs["Incoming"] == '"CREATE"'
    assert dvs[MERGE_ROW] == f'"{MERGE_ROW_ACTION}"'


def test_blank_merge_row_is_not_a_merge():
    """The blank default must read as "no decision", not as an empty merge."""
    payload = _conflict_payload()
    parsed = parse_curation_xlsx(build_curation_xlsx(payload))

    assert ("MERGE", KEY) not in parsed.field_actions


def test_clearing_the_incoming_action_folds_into_the_kept_row():
    """Blank on the incoming row is a deliberate choice, distinct from a group
    the sheet never mentioned -- which still defaults to CREATE."""
    payload = _conflict_payload()
    data = _edit(build_curation_xlsx(payload), INCOMING_ROW, action=None)
    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert plan.ok
    assert [(a.op, a.existing_id) for a in plan.actions] == [("reuse", 42)]


def test_no_dropdown_entry_is_blank():
    """A validation list cannot carry a blank entry - empty and whitespace-only
    items are both dropped when the dropdown is built, so offering one produces a
    menu that silently lost an option. Clearing the cell is the gesture for "no",
    which is why the optional actions set allow_blank instead.
    """
    data = build_curation_xlsx(_ambiguous_payload())
    ws = load_workbook(io.BytesIO(data))[SHEET]
    action_col = curation_headers().index("action") + 1

    for dv in ws.data_validations.dataValidation:
        first = ws[str(dv.sqref).split()[0].split(":")[0]]
        if first.column != action_col:
            continue
        options = dv.formula1.strip('"').split(",")
        assert all(o.strip() for o in options), dv.formula1
        # the optional ones must still accept a cleared cell
        optional = ws.cell(first.row, 1).value in (ROW_MARKER_INCOMING, MERGE_ROW)
        assert dv.allow_blank is optional


def test_merge_row_action_defaults_to_blank():
    ws = load_workbook(io.BytesIO(build_curation_xlsx(_conflict_payload())))[SHEET]
    action_col = curation_headers().index("action") + 1
    merge_row = next(
        r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == MERGE_ROW
    )

    assert ws.cell(merge_row, action_col).value in (None, "")


def _two_conflict_payload():
    """Two independent conflict groups, so one can be settled and one botched."""
    out = []
    for i in (0, 1):
        out.append(
            {
                "status": "conflict",
                "inchi_key": f"K{i}",
                "incoming": {"smiles": f"C{i}", "compound_code": f"NEW-{i}"},
                "conflicts": {
                    "compound_code": {"existing": f"OLD-{i}", "incoming": f"NEW-{i}"}
                },
                "existing": [
                    {
                        "id": 100 + i,
                        "inchi_key": f"K{i}",
                        "smiles": f"C{i}",
                        "compound_code": f"OLD-{i}",
                    }
                ],
            }
        )
    return out


def _merge_rows(data):
    ws = load_workbook(io.BytesIO(data))[SHEET]
    action_col = curation_headers().index("action") + 1
    code_col = curation_headers().index("compound_code") + 1
    return [
        (ws.cell(r, action_col).value, ws.cell(r, code_col))
        for r in range(1, ws.max_row + 1)
        if ws.cell(r, 1).value == MERGE_ROW
    ]


def test_reissued_workbook_carries_every_group_not_just_the_rejected_ones():
    """A group left out of the sheet reaches the resolver with no decision and
    silently falls back to the defaults -- discarding what the user already
    said. So the re-issued workbook must always be the full set."""
    payload = _two_conflict_payload()
    # settle the first group, botch the second (MERGE with nothing chosen)
    data = _edit(build_curation_xlsx(payload), MERGE_ROW, action=MERGE_ROW_ACTION)
    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET]
    code_col = curation_headers().index("compound_code") + 1
    first_merge = next(
        r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == MERGE_ROW
    )
    ws.cell(first_merge, code_col, "OLD-0")
    buf = io.BytesIO()
    wb.save(buf)

    parsed = parse_curation_xlsx(buf.getvalue())
    plan = resolve_curation(payload, parsed)
    assert len(plan.unresolved) == 1

    errors = {u.key: list(u.fields) for u in plan.unresolved if u.key}
    reissued = build_curation_xlsx(
        payload, error_cells=errors, prefill=parsed.field_actions
    )

    ws2 = load_workbook(io.BytesIO(reissued))[SHEET]
    assert sum(1 for r in ws2.iter_rows() if r[0].value == "id") == 2


def test_reissued_workbook_keeps_the_decisions_already_made():
    payload = _two_conflict_payload()
    data = _edit(build_curation_xlsx(payload), MERGE_ROW, action=MERGE_ROW_ACTION)
    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET]
    code_col = curation_headers().index("compound_code") + 1
    merges = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == MERGE_ROW]
    ws.cell(merges[0], code_col, "OLD-0")  # settled
    buf = io.BytesIO()
    wb.save(buf)

    parsed = parse_curation_xlsx(buf.getvalue())
    plan = resolve_curation(payload, parsed)
    errors = {u.key: list(u.fields) for u in plan.unresolved if u.key}
    rows = _merge_rows(
        build_curation_xlsx(payload, error_cells=errors, prefill=parsed.field_actions)
    )

    settled_action, settled_code = rows[0]
    botched_action, botched_code = rows[1]

    assert (settled_action, settled_code.value) == (MERGE_ROW_ACTION, "OLD-0")
    assert settled_code.fill.start_color.rgb == YELLOW  # accepted, not flagged
    # a MERGE whose values were all left blank parses to {} -- it must still come
    # back as MERGE, or the user's choice of merging is lost as well
    assert botched_action == MERGE_ROW_ACTION
    assert botched_code.value in (None, "")
    assert botched_code.fill.start_color.rgb == RED


def test_fixing_only_the_flagged_cells_settles_the_upload():
    """The round trip end to end: everything already accepted stays accepted."""
    payload = _two_conflict_payload()
    data = _edit(build_curation_xlsx(payload), MERGE_ROW, action=MERGE_ROW_ACTION)
    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET]
    code_col = curation_headers().index("compound_code") + 1
    merges = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == MERGE_ROW]
    ws.cell(merges[0], code_col, "OLD-0")
    buf = io.BytesIO()
    wb.save(buf)

    parsed = parse_curation_xlsx(buf.getvalue())
    plan = resolve_curation(payload, parsed)
    errors = {u.key: list(u.fields) for u in plan.unresolved if u.key}
    reissued = build_curation_xlsx(
        payload, error_cells=errors, prefill=parsed.field_actions
    )

    # the user fills in only what is red and sends it straight back
    wb2 = load_workbook(io.BytesIO(reissued))
    ws2 = wb2[SHEET]
    for r in range(1, ws2.max_row + 1):
        cell = ws2.cell(r, code_col)
        if cell.fill.start_color.rgb == RED:
            cell.value = "OLD-1"
    buf2 = io.BytesIO()
    wb2.save(buf2)

    final = resolve_curation(payload, parse_curation_xlsx(buf2.getvalue()))

    assert final.ok, [u.reason for u in final.unresolved]
    # both merges applied -- the first one was NOT reset to the defaults
    assert [(a.op, a.existing_id) for a in final.actions] == [
        ("update", 100),
        ("update", 101),
    ]


def _all_rows_set(payload, existing_action, incoming_action):
    data = build_curation_xlsx(payload)
    for ex_id in existing_row_ids(data):
        data = _edit(data, ex_id, action=existing_action)
    return _edit(data, INCOMING_ROW, action=incoming_action)


def test_deleting_everything_with_nothing_to_replace_it_is_refused():
    """Read literally the block says "delete all of these and add nothing". It
    must be refused, and the reason must name that, rather than talking about
    folding into a compound that would not exist."""
    payload = _ambiguous_payload()
    data = _all_rows_set(payload, "DELETE", None)

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert not plan.ok
    assert not plan.actions  # nothing is deleted: no action is produced at all
    (reason,) = [u.reason for u in plan.unresolved]
    assert "leave no compound at all" in reason


def test_deleting_everything_is_fine_when_the_incoming_one_replaces_them():
    """The same DELETEs are allowed once something survives to fold them into."""
    payload = _ambiguous_payload()
    data = _all_rows_set(payload, "DELETE", "CREATE")

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert plan.ok
    (action,) = plan.actions
    assert action.op == "create"
    assert sorted(action.superseded_ids) == [39, 40]


def test_too_many_survivors_reports_the_other_reason():
    payload = _ambiguous_payload()
    data = _all_rows_set(payload, "KEEP", None)

    plan = resolve_curation(payload, parse_curation_xlsx(data))

    assert not plan.ok
    (reason,) = [u.reason for u in plan.unresolved]
    assert "2 of them are marked KEEP" in reason
    assert "leave no compound" not in reason
