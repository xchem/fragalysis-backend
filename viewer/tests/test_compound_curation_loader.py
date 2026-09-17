"""Loader-side compound-curation tests: the reconciliation gate and the
resolution map it builds for process_compound.

The full create/reuse apply runs through the create_objects machinery and is
exercised by the loader integration test; here we cover the gate's decisions
(fail-safe on unresolved, and the reuse/update action map when the user has
supplied a completed curation spreadsheet).
"""

# pylint: disable=protected-access,redefined-outer-name,unused-argument

import io
import logging

import pytest
from openpyxl import load_workbook

from viewer.compound_curation import (
    ResolvedAction,
    auto_merge_non_conflicting_compounds,
    build_curation_xlsx,
    needs_curation,
)
from viewer.compound_reconciliation import inchi_key_for_smiles, reconcile_compounds
from viewer.models import Compound, Project, SiteObservation
from viewer.target_loader import (
    MetadataObject,
    ProcessedObject,
    TargetLoader,
    _row_named_by_pk,
)
from viewer.tests.curation_sheet_helpers import (
    MERGE_ROW,
    edit_group_row,
    existing_row_ids,
)

ETHANOL = "CCO"
CRYSTALS = {
    "Xtal-1": {
        "crystallographic_files": {
            "ligand_cif": {
                "ligands": {"LIG": {"smiles": ETHANOL, "compound_code": "NEW"}}
            }
        }
    }
}


class _Report:
    def __init__(self):
        self.logs = []

    def log(self, level, msg):
        self.logs.append((level, msg))

    def errors(self):
        return [m for lvl, m in self.logs if lvl == logging.ERROR]


def _loader(project, curation_file=None):
    tl = TargetLoader.__new__(TargetLoader)
    tl.project = project
    tl.report = _Report()  # type: ignore[assignment]
    tl.curation_file = curation_file
    tl._compound_resolution = {}
    tl._compound_supersede = {}
    tl._compound_keys = {}
    tl._compound_pk_watermark = 0
    tl.target = None
    return tl


def _existing(project, code="OLD"):
    return Compound.objects.create(
        project=project,
        smiles=ETHANOL,
        inchi="InChI=CCO",
        inchi_key=inchi_key_for_smiles(ETHANOL),
        compound_code=code,
    )


def test_extract_incoming_compounds_pure():
    assert TargetLoader._extract_incoming_compounds(CRYSTALS) == [
        {"smiles": ETHANOL, "compound_code": "NEW", "crystals": ["Xtal-1"]}
    ]


def test_extract_incoming_compounds_collects_every_crystal():
    """One compound in several crystals stays one entry and lists them all.

    The crystals are what the curation sheet shows, so a curator can trace a
    flagged compound back to its source data. Collecting them must not change
    what counts as a duplicate: the dedup key is the compound fields alone.
    """
    crystals = {
        "Xtal-2": CRYSTALS["Xtal-1"],
        "Xtal-1": CRYSTALS["Xtal-1"],
        "Xtal-3": {
            "crystallographic_files": {
                "ligand_cif": {
                    "ligands": {"LIG": {"smiles": "CCN", "compound_code": "OTHER"}}
                }
            }
        },
    }

    result = TargetLoader._extract_incoming_compounds(crystals)

    assert result == [
        {
            "smiles": ETHANOL,
            "compound_code": "NEW",
            "crystals": ["Xtal-1", "Xtal-2"],
        },
        {"smiles": "CCN", "compound_code": "OTHER", "crystals": ["Xtal-3"]},
    ]


@pytest.mark.django_db
def test_gate_fails_safe_on_unresolved_conflict():
    project = Project.objects.create(title="lb-1")
    _existing(project)
    tl = _loader(project)

    tl._reconcile_compounds_gate(CRYSTALS)

    assert tl.report.errors()  # ERROR logged -> triggers rollback
    assert tl._compound_resolution == {}  # nothing to reuse without a decision


@pytest.mark.django_db
def test_gate_builds_update_action_from_curation(tmp_path):
    project = Project.objects.create(title="lb-1")
    existing = _existing(project)

    recon = reconcile_compounds(project, [{"smiles": ETHANOL, "compound_code": "NEW"}])
    data = build_curation_xlsx(recon.curation_payload(), target_name="Mpro")

    # user takes the incoming values over the existing row: MERGE on the
    # existing duplicate, with the incoming compound_code picked in the merge row
    existing_id = existing_row_ids(data)[0]
    data = edit_group_row(data, existing_id, action="MERGE")
    data = edit_group_row(data, MERGE_ROW, action="MERGE", compound_code="NEW")

    path = tmp_path / "curation.xlsx"
    path.write_bytes(data)

    tl = _loader(project, curation_file=str(path))
    tl._reconcile_compounds_gate(CRYSTALS)

    assert not tl.report.errors()
    actions = list(tl._compound_resolution.values())
    assert len(actions) == 1
    assert actions[0].op == "update"
    assert actions[0].existing_id == existing.id
    assert actions[0].incoming["compound_code"] == "NEW"


@pytest.mark.django_db
def test_gate_reuse_action_for_exact_match():
    project = Project.objects.create(title="lb-1")
    existing = _existing(project, code="NEW")  # identical -> auto reuse, no conflict
    tl = _loader(project)

    tl._reconcile_compounds_gate(CRYSTALS)

    assert not tl.report.errors()
    actions = list(tl._compound_resolution.values())
    assert len(actions) == 1
    assert actions[0].op == "reuse"
    assert actions[0].existing_id == existing.id


@pytest.mark.django_db
def test_gate_noop_without_incoming():
    project = Project.objects.create(title="lb-1")
    tl = _loader(project)
    tl._reconcile_compounds_gate({})
    assert not tl.report.errors()
    assert tl._compound_resolution == {}


@pytest.mark.django_db
def test_reuse_applies_to_a_compound_no_target_links_to():
    """The bug this guards. The reachability check used
    Compound.filter_manager.by_target(), which finds compounds only through
    their site observations. A compound matched in the project but not yet
    linked to THIS target was therefore invisible, and a KEEP meaning "reuse
    compound N" silently created a duplicate instead - 481 of them in one real
    upload.
    """
    project = Project.objects.create(title="lb-1")
    existing = _existing(project)
    tl = _loader(project)

    # nothing links it to any target, which is precisely the blind spot
    assert not SiteObservation.objects.filter(cmpd=existing).exists()
    assert tl._resolved_compound_applies(existing.pk)


@pytest.mark.django_db
def test_reuse_does_not_reach_outside_the_project():
    """Widening from target to project must not widen further: reconciliation
    matched within one project, so a decision can only ever apply there."""
    project = Project.objects.create(title="lb-1")
    stranger = _existing(Project.objects.create(title="lb-2"))
    tl = _loader(project)

    assert not tl._resolved_compound_applies(stranger.pk)


@pytest.mark.django_db
def test_reuse_action_is_not_downgraded_for_an_unlinked_compound():
    """End of the chain: gate produces a reuse action, and the loader still
    considers it applicable when the compound belongs to no target."""
    project = Project.objects.create(title="lb-1")
    existing = _existing(project, code="NEW")  # exact match -> reuse, no curation
    tl = _loader(project)

    tl._reconcile_compounds_gate(CRYSTALS)

    assert not tl.report.errors()
    actions = list(tl._compound_resolution.values())
    assert [a.op for a in actions] == ["reuse"]
    assert actions[0].existing_id == existing.pk
    assert tl._resolved_compound_applies(actions[0].existing_id)


def _processed(fields):
    return ProcessedObject(
        model_class=Compound, fields=fields, key=("Xtal-1", "LIG"), defaults={}
    )


@pytest.mark.django_db
def test_pk_lookup_finds_a_row_the_target_scoped_manager_cannot_see():
    """Why this matters: create_objects looks instances up through
    by_target(), and on a miss it used to construct the model with whatever
    `fields` held and save it. With fields={"id": N} that is not an INSERT, it
    is an UPDATE of row N that blanks every column absent from `defaults` -- so
    reusing a sibling target's compound wiped the row it was meant to reuse.
    """
    project = Project.objects.create(title="lb-1")
    existing = _existing(project)

    assert not SiteObservation.objects.filter(cmpd=existing).exists()
    assert _row_named_by_pk(_processed({"id": existing.pk})) == existing


@pytest.mark.django_db
def test_pk_lookup_ignores_natural_key_fields():
    """Only an explicit primary key names a row globally; anything else is a
    normal lookup and must still be allowed to fall through to creation."""
    project = Project.objects.create(title="lb-1")
    _existing(project)

    assert _row_named_by_pk(_processed({"compound_code": "OLD"})) is None


@pytest.mark.django_db
def test_pk_lookup_returns_none_when_the_row_really_is_gone():
    project = Project.objects.create(title="lb-1")
    existing = _existing(project)
    missing = existing.pk
    existing.delete()

    assert _row_named_by_pk(_processed({"id": missing})) is None


# --------------------------------------------------------------------------- #
# Within-load deduplication
# --------------------------------------------------------------------------- #


class _Meta:
    """Stand-in for MetadataObject: process_compound only reads ``.new``."""

    def __init__(self, new=True):
        self.new = new


def _process(tl, experiments, item):
    """Call process_compound's undecorated body.

    ``@create_objects`` wraps it into a whole-yaml-block driver; what is under
    test is the single-item decision it returns.
    """
    # functools.wraps keeps the undecorated body here; pylint cannot see it.
    inner = TargetLoader.process_compound.__wrapped__  # pylint: disable=no-member
    return inner(tl, experiments=experiments, item_data=item)


def _item(crystal, smiles, compound_code):
    """An item_data tuple shaped as the create_objects flattener produces it."""
    return (
        crystal,
        "crystallographic_files",
        "ligand_cif",
        "ligands",
        "LIG",
        {"smiles": smiles, "compound_code": compound_code},
    )


def test_same_compound_in_two_crystals_reuses_one_row(db, make_project):
    """A second crystal carrying the same compound must not make a second row.

    The gate reconciles against the database as it stood *before* the load, so a
    compound this load has just created is invisible to it. Left alone, the same
    molecule soaked into N crystals produced N identical Compound rows - and
    every later upload was then permanently ambiguous, because reconciliation
    cannot choose between identical rows. Measured on a real six-upload target:
    one clean load of A71EV2A produced 37 duplicated InChI keys.
    """
    project = make_project("proposal")
    tl = _loader(project)
    experiments = {"Xtal-1": _Meta(), "Xtal-2": _Meta()}

    first = _process(tl, experiments, _item("Xtal-1", ETHANOL, "CODE-1"))
    # Nothing to reuse yet, so the loader is told to create.
    assert first.fields == {}

    # Stand in for create_objects having created that row.
    created = Compound.objects.create(**first.defaults)

    second = _process(tl, experiments, _item("Xtal-2", ETHANOL, "CODE-1"))
    assert second.fields == {"id": created.pk}
    assert _row_named_by_pk(second) == created


def test_a_different_compound_code_is_left_for_curation(db, make_project):
    """The same structure under another code is a curation decision, not a dupe.

    ``viewer.compound_dedup`` groups exact duplicates by
    (project, inchi_key, smiles, compound_code); anything short of that carries
    information one row does not have, so the loader must not silently collapse
    it.
    """
    project = make_project("proposal")
    tl = _loader(project)
    experiments = {"Xtal-1": _Meta(), "Xtal-2": _Meta()}

    first = _process(tl, experiments, _item("Xtal-1", ETHANOL, "CODE-1"))
    Compound.objects.create(**first.defaults)

    second = _process(tl, experiments, _item("Xtal-2", ETHANOL, "CODE-2"))
    assert second.fields == {}


def test_compounds_from_earlier_loads_are_left_to_the_gate(db, make_project):
    """Only rows created by THIS load are reused; older ones stay the gate's job.

    An ambiguity that predates the load is exactly what the curation gate exists
    to resolve, and quietly reusing one of the candidates here would pre-empt a
    decision the user is supposed to make.
    """
    project = make_project("proposal")
    tl = _loader(project)
    experiments = {"Xtal-1": _Meta()}

    probe = _process(tl, experiments, _item("Xtal-1", ETHANOL, "CODE-1"))
    earlier = Compound.objects.create(**probe.defaults)
    # The gate ran after that row existed.
    tl._compound_pk_watermark = earlier.pk

    again = _process(tl, experiments, _item("Xtal-1", ETHANOL, "CODE-1"))
    assert again.fields == {}


@pytest.mark.django_db
def test_validation_chain_survives_an_uploader_that_sends_no_crystals():
    """The exact sequence UploadTargetExperimentsValidate runs, on a legacy payload.

    Fragalysis and XCA are released separately, so a backend carrying the
    crystal column will be asked to reconcile payloads that predate it. The
    whole chain - reconcile, auto-merge, filter, render - has to come through
    with the column simply blank.
    """
    project = Project.objects.create(title="lb-1")
    _existing(project)

    # No "crystals" key: what an older uploader sends.
    compounds = [{"smiles": ETHANOL, "compound_code": "NEW"}]

    reconciliation = reconcile_compounds(project, compounds)
    curation = reconciliation.curation_payload()
    curation, _stats = auto_merge_non_conflicting_compounds(curation)
    curation = needs_curation(curation)
    assert curation, "this payload should still need a decision"

    data = build_curation_xlsx(curation, target_name="Mpro")

    ws = load_workbook(io.BytesIO(data))["Incoming compounds"]
    cols = {}
    blanks = 0
    for row in ws.iter_rows():
        if row[0].value == "id":
            cols = {c.value: c.column for c in row if c.value}
            continue
        if row[0].value is not None:
            assert ws.cell(row[0].row, cols["crystal"]).value in (None, "")
            blanks += 1
    assert blanks  # rows were actually written


@pytest.mark.django_db
def test_supersession_repoints_cached_instances_at_the_survivor():
    """A retired compound must not stay cached in compound_objects.

    process_compound runs before the supersessions, so its output can hold an
    instance of a row the curation then deletes. Everything downstream reads its
    compound from that dict - the experiment.compounds link loop immediately
    after, and process_site_observation later - so a stale instance is written
    back as a reference to a compound that no longer exists. Postgres defers the
    check, so it surfaces at COMMIT as a bare foreign-key violation with nothing
    to tie it to the upload:

        update or delete on table "viewer_compound" violates foreign key
        constraint ... on table "viewer_experimentcompound"

    Only reachable through MERGE or DELETE: CREATE retires nothing.
    """
    project = Project.objects.create(title="lb-1")
    survivor = _existing(project, code="KEEP-ME")
    doomed = _existing(project, code="RETIRE-ME")
    doomed_pk = doomed.pk

    tl = _loader(project)
    tl._compound_supersede = {
        "row-1": ResolvedAction(
            index=0,
            op="update",
            existing_id=survivor.pk,
            incoming={},
            superseded_ids=[doomed_pk],
        )
    }
    cached = MetadataObject(instance=doomed, key="k", versioned_key="k")
    compound_objects = {("Xtal-1", "LIG"): cached}

    tl._apply_compound_supersessions(compound_objects)

    assert not Compound.objects.filter(pk=doomed_pk).exists()
    assert cached.instance.pk == survivor.pk
    # and the report names the row that went, not the None a deleted instance
    # reports for its pk
    assert str(doomed_pk) in " ".join(m for _lvl, m in tl.report.logs)
